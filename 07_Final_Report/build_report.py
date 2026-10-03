#!/usr/bin/env python3
"""
Assemble the IE3122 Northbridge report into a single Word document.

Run from the assignment root:
    python 07_Final_Report/build_report.py

Produces: IE3122-Northbridge-Report.docx
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
OUT = "IE3122-Northbridge-Report.docx"

# ---------------------------------------------------------------- styling

def setup_styles(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"
    n.font.size = Pt(11)
    n.paragraph_format.space_after = Pt(6)
    n.paragraph_format.line_spacing = 1.15

    for name, size, color, before in [
        ("Heading 1", 18, "1F3864", 18),
        ("Heading 2", 14, "2E5496", 14),
        ("Heading 3", 12, "2E5496", 10),
        ("Heading 4", 11, "404040", 8),
    ]:
        s = doc.styles[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True

    if "CodeBlock" not in [s.name for s in doc.styles]:
        from docx.enum.style import WD_STYLE_TYPE
        c = doc.styles.add_style("CodeBlock", WD_STYLE_TYPE.PARAGRAPH)
        c.font.name = "Consolas"
        c.font.size = Pt(8.5)
        c.paragraph_format.space_after = Pt(0)
        c.paragraph_format.space_before = Pt(0)
        c.paragraph_format.line_spacing = 1.0
        c.paragraph_format.left_indent = Inches(0.25)

        q = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
        q.font.name = "Calibri"
        q.font.size = Pt(10)
        q.font.italic = True
        q.font.color.rgb = RGBColor.from_string("404040")
        q.paragraph_format.left_indent = Inches(0.3)
        q.paragraph_format.space_after = Pt(6)

        cap = doc.styles.add_style("Caption2", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.name = "Calibri"
        cap.font.size = Pt(9)
        cap.font.italic = True
        cap.font.color.rgb = RGBColor.from_string("404040")
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(12)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def shade(cell, hexcolor):
    el = OxmlElement("w:shd")
    el.set(qn("w:fill"), hexcolor)
    cell._tc.get_or_add_tcPr().append(el)

# ---------------------------------------------------------------- inline md

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")

def add_runs(par, text):
    """Render **bold**, *italic*, `code` into a paragraph."""
    text = text.replace("\\|", "|")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("***") and part.endswith("***"):
            r = par.add_run(part[3:-3]); r.bold = True; r.italic = True
        elif part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            par.add_run(part)

def split_row(line):
    line = line.strip()
    if line.startswith("|"): line = line[1:]
    if line.endswith("|"): line = line[:-1]
    out, cur, i = [], "", 0
    while i < len(line):
        if line[i] == "\\" and i + 1 < len(line) and line[i+1] == "|":
            cur += "\\|"; i += 2; continue
        if line[i] == "|":
            out.append(cur.strip()); cur = ""; i += 1; continue
        cur += line[i]; i += 1
    out.append(cur.strip())
    return out

# ---------------------------------------------------------------- md -> docx

def render_markdown(doc, md, heading_offset=0, drop_h1=False):
    lines = md.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            if buf:
                for b in buf:
                    p = doc.add_paragraph(style="CodeBlock")
                    p.add_run(b if b.strip() else " ")
                doc.add_paragraph()
            continue

        # table
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|?$", lines[i+1].strip()):
            header = split_row(stripped)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip())); i += 1
            ncols = len(header)
            t = doc.add_table(rows=1, cols=ncols)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, h in enumerate(header):
                cell = t.rows[0].cells[c]
                cell.text = ""
                add_runs(cell.paragraphs[0], h)
                for r in cell.paragraphs[0].runs: r.bold = True
                shade(cell, "DEEAF6")
            for row in rows:
                cells = t.add_row().cells
                for c in range(ncols):
                    val = row[c] if c < len(row) else ""
                    cells[c].text = ""
                    add_runs(cells[c].paragraphs[0], val)
            for row in t.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(2)
                        for r in p.runs: r.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            lvl = len(m.group(1))
            txt = m.group(2).strip()
            if drop_h1 and lvl == 1:
                i += 1; continue
            lvl = min(lvl + heading_offset, 4)
            p = doc.add_paragraph(style=f"Heading {max(lvl,1)}")
            add_runs(p, txt)
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(
                    {1:"1F3864",2:"2E5496",3:"2E5496",4:"404040"}[max(lvl,1)])
            i += 1; continue

        # blockquote
        if stripped.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            sub = "\n".join(buf).strip()
            if sub:
                # a quote may itself contain a table
                if "|" in sub and re.search(r"^\|[\s:\-|]+\|?$", sub, re.M):
                    render_markdown(doc, sub, heading_offset, drop_h1)
                else:
                    for para in [x for x in sub.split("\n\n") if x.strip()]:
                        p = doc.add_paragraph(style="Note")
                        add_runs(p, " ".join(para.split("\n")))
            continue

        # horizontal rule
        if re.match(r"^(---|\*\*\*|___)\s*$", stripped):
            i += 1; continue

        # lists
        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            style = "List Number" if ordered else "List Bullet"
            p = doc.add_paragraph(style=style)
            if indent >= 2:
                p.paragraph_format.left_indent = Inches(0.5 + 0.25 * (indent // 2))
            add_runs(p, m.group(3))
            p.paragraph_format.space_after = Pt(3)
            i += 1; continue

        # blank
        if not stripped:
            i += 1; continue

        # paragraph (gather continuation lines)
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", ">", "|", "```", "---"))
                    or re.match(r"^(\s*)([-*+]|\d+\.)\s+", lines[i])):
                break
            buf.append(nxt); i += 1
        p = doc.add_paragraph()
        add_runs(p, " ".join(buf))

# ---------------------------------------------------------------- captions

CAPTIONS = {
 "1.1":"AAA server: service enabled, both protocol clients and two user accounts",
 "1.2":"AAA configuration entered on HQ-R1",
 "1.3":"Method lists in the running configuration, showing group tacacs+ local",
 "1.4":"TACACS+ and RADIUS server definitions",
 "1.5":"Active AAA session listed by user name",
 "1.6":"show aaa user all reporting method=TACACS for the login",
 "1.7":"SSH login by alice.perera, an account held only on the AAA server",
 "1.8":"Login rejected for an account that exists nowhere",
 "1.9":"Local fallback: AAA server unreachable, netadmin still admitted",
 "1.10":"RSA key pair generated",
 "1.11":"SSH enabled, version 2.0",
 "1.12":"Key pair HQ-R1.northbridge.lk at 1024 bits",
 "1.13":"VTY lines: login local, SSH only, access-class and exec-timeout",
 "1.14":"MGMT-HOSTS access list before testing, counters at zero",
 "1.15":"SSH from the management subnet succeeds, banner displayed",
 "1.16":"Telnet refused from an authorised host, proving the protocol is disabled",
 "1.17":"SSH refused from outside the management subnet, no password prompt",
 "1.18":"MGMT-HOSTS deny counter incremented by the blocked attempt",
 "2.1":"Six VLANs with correct port membership",
 "2.2":"Trunk with explicit allowed list and native VLAN moved to 999",
 "2.3":"Unused ports administratively disabled in the blackhole VLAN",
 "2.4":"Five subinterfaces up with their zone gateway addresses",
 "2.5":"One connected route per VLAN",
 "2.6":"All four inter-VLAN access lists, counters at zero",
 "2.7":"Permitted business flow: staff reach the core banking application over HTTPS",
 "2.8":"Staff denied ICMP to core banking and all access to management",
 "2.9":"Guest Wi-Fi reaches neither core banking nor staff",
 "2.10":"Access list counters after testing: permitted and denied traffic both recorded",
 "2.11":"Port security summary with per-role violation actions",
 "2.12":"Fa0/1 secure-up, maximum one address, violation mode shutdown",
 "2.13":"Secure MAC address table with sticky entries",
 "2.14":"Learned MAC written into the running configuration",
 "2.14b":"Running configuration, continued",
 "2.14c":"Running configuration, continued",
 "2.15":"Unauthorised PC connected in place of the authorised workstation",
 "2.16":"Port in secure-shutdown with violation count 1 and the offending MAC recorded",
 "2.17":"Interface err-disabled",
 "2.18":"Violation counted in the port security summary",
 "2.19":"Port recovered to secure-up after the authorised device was restored",
 "3.1":"Zone-Based Policy Firewall rejected by the platform",
 "3.1b":"CBAC and reflexive ACLs also rejected; no inspect keyword in the parser",
 "3.2":"securityk9 still disabled after reload, confirming the licence command is a no-op",
 "3.3":"Edge firewall configuration entered",
 "3.4":"NAT boundary: outside interface and four inside subinterfaces",
 "3.5":"Static PAT publish and outbound overload rules",
 "3.6":"EDGE-IN access list before testing, counters at zero",
 "3.7":"Online banking portal reachable from the internet on its public address",
 "3.7a":"Portal timing out before the established fix, demonstrating the stateless-ACL limitation",
 "3.8":"Internet host denied access to internal staff and core banking",
 "3.9":"EDGE-IN counters after testing: published service permitted, internal access denied",
 "3.10":"Live NAT translation carrying the portal session",
 "3.11":"Outbound PAT hiding internal addressing",
 "3.12":"DHCP snooping enabled, uplink the only trusted interface",
 "3.13":"Binding table built from observed DHCP exchanges",
 "3.14":"Dynamic ARP Inspection active on the client VLANs",
 "3.15":"Per-interface trust state and rate limits",
 "3.16":"Per-VLAN inspection statistics",
 "3.17":"Legitimate DHCP working before the attack",
 "3.18":"Rogue DHCP server connected to a staff access port",
 "3.19":"Victim retains the legitimate gateway and DNS; the rogue offer was dropped",
 "3.20":"Binding table after the attack contains no rogue-sourced entry",
 "3.20b":"The rogue server, live and correctly configured to advertise itself as gateway",
 "3.21":"DAI enabled and active (Packet Tracer does not increment its counters)",
 "4.1":"Crypto command family absent from the parser",
 "4.1b":"No licence activation path available",
 "4.1c":"securityk9 still disabled after two reloads",
 "4.2":"NetFlow export configuration entered",
 "4.3":"Flow cache operational before traffic",
 "4.4":"Flow collection and export statements in the running configuration",
 "4.5":"Flow records for normal traffic",
 "4.6a":"Reconnaissance sequence generated",
 "4.6b":"Reconnaissance sequence, continued",
 "4.6c":"Scan attempts, several refused by the inter-VLAN access lists",
 "4.7":"Flow cache aggregate: 153 ICMP flows recorded by the scan",
 "4.7b":"Fan-out: one source address reaching multiple destinations across subnets",
 "4.7c":"Flow cache timeout not configurable in Packet Tracer",
 "4.8":"Branch flow visibility with no appliance deployed at the branch",
 "4.10":"Head-office clock set, time source NTP",
 "4.11":"Logging to 10.10.99.11 active, messages sent",
 "4.12":"Branch switch synchronised at stratum 4",
 "4.12b":"Head-office switch synchronised",
 "4.13":"NTP associations with the selected peer",
 "4.14":"ntp source rejected, matching the logging source-interface limitation",
 "4.16":"Messages arriving at the collector with millisecond timestamps and source device",
 "4.18":"Failed login recorded centrally",
 "4.19":"Events from two devices at two sites on one ordered timeline",
}

def caption_for(fname):
    m = re.match(r"^(\d+\.\d+[a-z]?)", fname)
    key = m.group(1) if m else None
    txt = CAPTIONS.get(key)
    if not txt:
        txt = re.sub(r"^[\d.]+[a-z]?-", "", os.path.splitext(fname)[0]).replace("-", " ").capitalize()
    return key, txt

def shot_sort_key(fname):
    m = re.match(r"^(\d+)\.(\d+)([a-z]?)", fname)
    if not m: return (99, 99, "")
    return (int(m.group(1)), int(m.group(2)), m.group(3))

FIG = {"n": 0}

def add_screenshots(doc, folder, title):
    d = os.path.join("08_Screenshots", folder)
    if not os.path.isdir(d): return 0
    files = sorted([f for f in os.listdir(d) if f.lower().endswith(".png")], key=shot_sort_key)
    if not files: return 0
    h = doc.add_paragraph(style="Heading 4"); add_runs(h, title)
    for f in files:
        key, cap = caption_for(f)
        try:
            doc.add_picture(os.path.join(d, f), width=Inches(6.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        except Exception as e:
            p = doc.add_paragraph(); p.add_run(f"[image could not be embedded: {f} — {e}]").italic = True
            continue
        FIG["n"] += 1
        c = doc.add_paragraph(style="Caption2")
        add_runs(c, f"Figure {FIG['n']} — Shot {key}: {cap}")
    return len(files)

# ---------------------------------------------------------------- helpers

def read(p):
    with open(p, encoding="utf-8") as f: return f.read()

def h1(doc, text, page_break=True):
    if page_break:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(style="Heading 1"); add_runs(p, text)
    return p

def script_block(doc, path, title):
    p = doc.add_paragraph(style="Heading 3"); add_runs(p, title)
    for line in read(path).split("\n"):
        par = doc.add_paragraph(style="CodeBlock")
        par.add_run(line if line.strip() else " ")
    doc.add_paragraph()

# ---------------------------------------------------------------- build

def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.9)
        s.top_margin = s.bottom_margin = Inches(0.8)
    setup_styles(doc)

    # ---- cover
    render_markdown(doc, read("07_Final_Report/00-cover-page.md"))

    # ---- toc placeholder
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph(style="Heading 1"); add_runs(p, "Table of Contents")
    p = doc.add_paragraph(style="Note")
    add_runs(p, "In Word: References → Table of Contents → Automatic Table. "
                "It will populate from the heading styles used throughout this document.")

    # ---- 1 exec summary
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    render_markdown(doc, read("07_Final_Report/02-executive-summary.md"))

    # ---- 2 & 3 scope + context
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    render_markdown(doc, read("07_Final_Report/03-scope-and-context.md"))

    # ---- 4 threat modelling
    h1(doc, "4. Individual Threat Modelling")
    p = doc.add_paragraph(style="Note")
    add_runs(p, "Eight threats, two per member, each individually authored and individually "
                "attributable as the brief requires. All four members used the common risk "
                "method set out in Section 2.3.")
    for n, folder in enumerate(["01_Member_1","02_Member_2","03_Member_3","04_Member_4"], 1):
        hp = doc.add_paragraph(style="Heading 2"); add_runs(hp, f"4.{n} Member {n}")
        render_markdown(doc, read(f"{folder}/01-threat-models.md"), heading_offset=2, drop_h1=True)

    # ---- 5 layered plan
    h1(doc, "5. Combined Layered Security Plan")
    render_markdown(doc, read("05_Group/01-layered-security-plan.md"), heading_offset=1, drop_h1=True)

    # ---- 6 tables
    h1(doc, "6. Control Mapping and Classification")
    render_markdown(doc, read("05_Group/02-control-mapping-tables.md"), heading_offset=1, drop_h1=True)

    # ---- 7 minimisation
    h1(doc, "7. Control Minimisation Justification")
    render_markdown(doc, read("05_Group/03-control-minimisation.md"), heading_offset=1, drop_h1=True)

    # ---- 8 configurations + screenshots
    h1(doc, "8. Individual Network Hardening Configurations")
    p = doc.add_paragraph(style="Note")
    add_runs(p, "Eight distinct configurations, two per member, at least one of each pair "
                "implementing a control from the plan. Screenshot evidence follows each member's "
                "write-up. Plain-text scripts are in Appendix B.")
    shots = {}
    plan = [
        ("01_Member_1", 1, [("M1-cfg1-aaa","Configuration 1 — AAA with TACACS+ and RADIUS"),
                            ("M1-cfg2-ssh-mgmt","Configuration 2 — SSH and management-plane hardening")]),
        ("02_Member_2", 2, [("M2-cfg3-vlan-acl","Configuration 3 — VLAN segmentation with inter-VLAN ACLs"),
                            ("M2-cfg4-port-security","Configuration 4 — Switchport port security")]),
        ("03_Member_3", 3, [("M3-cfg5-edge-firewall","Configuration 5 — Internet edge firewall"),
                            ("M3-cfg6-dhcp-snoop-dai","Configuration 6 — DHCP snooping and Dynamic ARP Inspection")]),
        ("04_Member_4", 4, [("M4-cfg7-netflow","Configuration 7 — NetFlow export"),
                            ("M4-cfg8-syslog-ntp","Configuration 8 — Centralised syslog and NTP")]),
    ]
    for folder, n, cfgs in plan:
        hp = doc.add_paragraph(style="Heading 2"); add_runs(hp, f"8.{n} Member {n}")
        render_markdown(doc, read(f"{folder}/03-configurations.md"), heading_offset=2, drop_h1=True)
        hp = doc.add_paragraph(style="Heading 3"); add_runs(hp, f"8.{n}.E  Screenshot evidence")
        for key, title in cfgs:
            shots[key] = add_screenshots(doc, key, title)

    # ---- baseline evidence
    hp = doc.add_paragraph(style="Heading 2"); add_runs(hp, "8.5 Pre-control baseline evidence")
    p = doc.add_paragraph()
    add_runs(p, "Captured on the verified topology **before any control was applied**. Each pairs "
                "with an 'after' shot where the same test is refused, demonstrating the control "
                "against a measured baseline rather than in isolation.")
    shots["00-BEFORE"] = add_screenshots(doc, "00-BEFORE", "Baseline — the network as Northbridge runs it today")

    # ---- 9 technology evaluation
    h1(doc, "9. Group Technology Evaluation and Selection")
    render_markdown(doc, read("05_Group/04-group-technology-evaluation.md"), heading_offset=1, drop_h1=True)

    # ---- 10 topology
    h1(doc, "10. Topology Used for the Practical Work")
    render_markdown(doc, read("06_Topology/topology-spec.md"), heading_offset=1, drop_h1=True)
    if os.path.exists("topology-diagram.png"):
        doc.add_picture("topology-diagram.png", width=Inches(6.5))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        FIG["n"] += 1
        c = doc.add_paragraph(style="Caption2")
        add_runs(c, f"Figure {FIG['n']} — The Packet Tracer topology as built")

    # ---- appendices A1-A4
    for n, folder in enumerate(["01_Member_1","02_Member_2","03_Member_3","04_Member_4"], 1):
        h1(doc, f"Appendix A{n} — Member {n}: Individual Technology Proposals")
        render_markdown(doc, read(f"{folder}/02-tech-proposals.md"), heading_offset=1, drop_h1=True)

    # ---- appendix B scripts
    h1(doc, "Appendix B — Configuration Scripts (plain text)")
    p = doc.add_paragraph(style="Note")
    add_runs(p, "Exactly as entered. Commands rejected by Packet Tracer are commented out with "
                "the reason and the production equivalent, as required by the brief.")
    for folder, titles in [
        ("01_Member_1", [("cfg-1-aaa-tacacs-radius.txt","B1 — Configuration 1: AAA with TACACS+ and RADIUS"),
                         ("cfg-2-ssh-mgmt-hardening.txt","B2 — Configuration 2: SSH and management-plane hardening")]),
        ("02_Member_2", [("cfg-3-vlan-segmentation-acl.txt","B3 — Configuration 3: VLAN segmentation with inter-VLAN ACLs"),
                         ("cfg-4-port-security.txt","B4 — Configuration 4: Switchport port security")]),
        ("03_Member_3", [("cfg-5-edge-firewall-nat.txt","B5 — Configuration 5: Internet edge firewall"),
                         ("cfg-6-dhcp-snooping-dai.txt","B6 — Configuration 6: DHCP snooping and Dynamic ARP Inspection")]),
        ("04_Member_4", [("cfg-7-netflow-export.txt","B7 — Configuration 7: NetFlow export"),
                         ("cfg-8-syslog-ntp.txt","B8 — Configuration 8: Centralised syslog and NTP")]),
    ]:
        for fn, title in titles:
            path = os.path.join(folder, "configs", fn)
            if os.path.exists(path):
                script_block(doc, path, title)

    doc.save(OUT)

    total = sum(shots.values())
    print(f"\n  Written: {OUT}")
    print(f"  Figures embedded: {FIG['n']}")
    for k, v in shots.items():
        print(f"    {v:3d}  {k}")
    print(f"    {total:3d}  TOTAL screenshots")

if __name__ == "__main__":
    main()
