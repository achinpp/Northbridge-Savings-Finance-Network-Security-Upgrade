#!/usr/bin/env python3
"""
Build the condensed IE3122 report, target under 40 pages.

Run from the assignment root:
    python 09_Concise/build_concise.py
"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
SRC = "09_Concise"
OUT = "IE3122-Northbridge-Report-Concise.docx"

BODY, CODE, CAP = 9.0, 6.5, 7.0
IMG_W = 2.95

# ------------------------------------------------------------- selection

SHOTS = {
 "00-BEFORE": ["guest-to-coredb", "internet-to-staff", "dmz-to-coredb"],
 "M1-cfg1-aaa": ["1.1-", "1.7-", "1.9-"],
 "M1-cfg2-ssh-mgmt": ["1.13-", "1.16-", "1.18-"],
 "M2-cfg3-vlan-acl": ["2.7-", "2.8-", "2.10-"],
 "M2-cfg4-port-security": ["2.11-", "2.16-", "2.19-"],
 "M3-cfg5-edge-firewall": ["3.1-", "3.7-p", "3.8-", "3.9-"],
 "M3-cfg6-dhcp-snoop-dai": ["3.13-", "3.19-", "3.20-b"],
 "M4-cfg7-netflow": ["4.1-", "4.5-", "4.7b-"],
 "M4-cfg8-syslog-ntp": ["4.12-n", "4.16-", "4.19-"],
}

GROUPS = [
 ("8.1 Member 1", [("M1-cfg1-aaa", "Configuration 1, AAA"),
                   ("M1-cfg2-ssh-mgmt", "Configuration 2, SSH and management plane")]),
 ("8.2 Member 2", [("M2-cfg3-vlan-acl", "Configuration 3, VLAN segmentation and ACLs"),
                   ("M2-cfg4-port-security", "Configuration 4, port security")]),
 ("8.3 Member 3", [("M3-cfg5-edge-firewall", "Configuration 5, internet edge firewall"),
                   ("M3-cfg6-dhcp-snoop-dai", "Configuration 6, DHCP snooping and DAI")]),
 ("8.4 Member 4", [("M4-cfg7-netflow", "Configuration 7, NetFlow export"),
                   ("M4-cfg8-syslog-ntp", "Configuration 8, syslog and NTP")]),
 ("8.5 Baseline", [("00-BEFORE", "Before any control was applied")]),
]

CAPTIONS = {
 "guest-to-coredb": "Guest Wi-Fi reaching the core banking database",
 "internet-to-staff": "An internet host reaching an internal staff workstation",
 "dmz-to-coredb": "The public web portal reaching the core banking database",
 "1.1": "AAA server, service on, both protocol clients, two users",
 "1.2": "AAA configuration entered on HQ-R1",
 "1.7": "SSH login by a user held only on the AAA server",
 "1.9": "Local fallback admits netadmin while the AAA server is unreachable",
 "1.11": "SSH enabled, version 2.0",
 "1.13": "VTY lines, login local, SSH only, access class, timeout",
 "1.16": "Telnet refused from an authorised host",
 "1.18": "MGMT-HOSTS deny counter incremented by the blocked attempt",
 "2.2": "Trunk with allowed list and native VLAN moved to 999",
 "2.6": "Four inter VLAN access lists, counters at zero",
 "2.7": "Staff reach core banking over HTTPS, the permitted flow",
 "2.8": "Staff denied ICMP to core banking and all access to management",
 "2.10": "Counters after testing, permitted and denied traffic both recorded",
 "2.11": "Port security summary with per role violation actions",
 "2.14": "Learned MAC written into the running configuration",
 "2.16": "Secure-shutdown, violation count 1, offending MAC recorded",
 "2.19": "Port recovered to Secure-up",
 "3.1": "Zone Based Policy Firewall rejected by the platform",
 "3.2": "securityk9 still disabled after reload",
 "3.5": "Static PAT publish and outbound overload rules",
 "3.7": "Portal reachable from the internet on its public address",
 "3.8": "Internet denied access to internal staff and core banking",
 "3.9": "EDGE-IN counters, published service permitted, internal denied",
 "3.12": "DHCP snooping enabled, uplink the only trusted interface",
 "3.13": "Binding table built from observed DHCP exchanges",
 "3.18": "Rogue DHCP server connected to a staff access port",
 "3.19": "Victim keeps the legitimate gateway and DNS, rogue offer dropped",
 "3.20": "The rogue server, live and correctly configured",
 "4.1": "Crypto command family absent from the parser",
 "4.3": "Flow cache operational before traffic",
 "4.5": "Flow records for normal traffic",
 "4.7b": "Fan out, one source reaching multiple destinations",
 "4.8": "Branch flow visibility with no appliance at the branch",
 "4.11": "Logging to 10.10.99.11 active, messages sent",
 "4.12": "Branch switch synchronised at stratum 4",
 "4.14": "ntp source rejected, matching the logging limitation",
 "4.16": "Messages arriving with timestamps and source device",
 "4.19": "Two devices at two sites on one ordered timeline",
}

FIG = {"n": 0}

# ------------------------------------------------------------- styling

def setup(doc):
    n = doc.styles["Normal"]
    n.font.name = "Calibri"; n.font.size = Pt(BODY)
    n.paragraph_format.space_after = Pt(2)
    n.paragraph_format.line_spacing = 1.0

    for name, size, col, before in [("Heading 1", 14, "1F3864", 10),
                                    ("Heading 2", 11.5, "2E5496", 8),
                                    ("Heading 3", 10, "2E5496", 6),
                                    ("Heading 4", 9.5, "404040", 5)]:
        s = doc.styles[name]
        s.font.name = "Calibri"; s.font.size = Pt(size); s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(col)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(2)
        s.paragraph_format.keep_with_next = True

    c = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    c.font.name = "Consolas"; c.font.size = Pt(CODE)
    c.paragraph_format.space_after = Pt(0); c.paragraph_format.space_before = Pt(0)
    c.paragraph_format.line_spacing = 0.95
    c.paragraph_format.left_indent = Inches(0.15)

    q = doc.styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    q.font.name = "Calibri"; q.font.size = Pt(BODY - 0.5); q.font.italic = True
    q.font.color.rgb = RGBColor.from_string("404040")
    q.paragraph_format.left_indent = Inches(0.2)
    q.paragraph_format.space_after = Pt(3)

    cp = doc.styles.add_style("Cap", WD_STYLE_TYPE.PARAGRAPH)
    cp.font.name = "Calibri"; cp.font.size = Pt(CAP); cp.font.italic = True
    cp.font.color.rgb = RGBColor.from_string("404040")
    cp.paragraph_format.space_before = Pt(1); cp.paragraph_format.space_after = Pt(5)
    cp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def shade(cell, hexc):
    el = OxmlElement("w:shd"); el.set(qn("w:fill"), hexc)
    cell._tc.get_or_add_tcPr().append(el)

def no_borders(table):
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'none'); e.set(qn('w:sz'), '0')
        borders.append(e)
    pr.append(borders)

# ------------------------------------------------------------- md render

INLINE = re.compile(r"(\*\*\*.+?\*\*\*|\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`)")

def runs(par, text):
    text = text.replace("\\|", "|")
    for part in INLINE.split(text):
        if not part: continue
        if part.startswith("***") and part.endswith("***"):
            r = par.add_run(part[3:-3]); r.bold = True; r.italic = True
        elif part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(BODY - 1.5)
        elif part.startswith("*") and part.endswith("*"):
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            par.add_run(part)

def cells_of(line):
    line = line.strip().strip("|")
    out, cur, i = [], "", 0
    while i < len(line):
        if line[i] == "\\" and i+1 < len(line) and line[i+1] == "|":
            cur += "\\|"; i += 2; continue
        if line[i] == "|":
            out.append(cur.strip()); cur = ""; i += 1; continue
        cur += line[i]; i += 1
    out.append(cur.strip()); return out

def render(doc, md, off=0, drop_h1=False):
    lines = md.split("\n"); i = 0
    while i < len(lines):
        line = lines[i]; s = line.strip()

        if s.startswith("```"):
            i += 1; buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph(style="Code"); p.add_run(b if b.strip() else " ")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        if s.startswith("|") and i+1 < len(lines) and re.match(r"^\|[\s:\-|]+\|?$", lines[i+1].strip()):
            hdr = cells_of(s); i += 2; rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(cells_of(lines[i])); i += 1
            nc = len(hdr)
            t = doc.add_table(rows=1, cols=nc); t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, h in enumerate(hdr):
                cl = t.rows[0].cells[c]; cl.text = ""
                runs(cl.paragraphs[0], h)
                for r in cl.paragraphs[0].runs: r.bold = True
                shade(cl, "DEEAF6")
            for row in rows:
                cs = t.add_row().cells
                for c in range(nc):
                    cs[c].text = ""; runs(cs[c].paragraphs[0], row[c] if c < len(row) else "")
            for row in t.rows:
                for cl in row.cells:
                    for p in cl.paragraphs:
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.line_spacing = 1.0
                        for r in p.runs: r.font.size = Pt(7.0)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = len(m.group(1))
            if drop_h1 and lvl == 1: i += 1; continue
            lvl = min(lvl + off, 4)
            p = doc.add_paragraph(style=f"Heading {max(lvl,1)}"); runs(p, m.group(2).strip())
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor.from_string({1:"1F3864",2:"2E5496",3:"2E5496",4:"404040"}[max(lvl,1)])
            i += 1; continue

        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i])); i += 1
            sub = "\n".join(buf).strip()
            if sub:
                if re.search(r"^\|[\s:\-|]+\|?$", sub, re.M):
                    render(doc, sub, off, drop_h1)
                else:
                    for para in [x for x in sub.split("\n\n") if x.strip()]:
                        p = doc.add_paragraph(style="Note"); runs(p, " ".join(para.split("\n")))
            continue

        m = re.match(r"^!\[(.*?)\]\((.+?)\)\s*$", s)
        if m:
            alt, path = m.group(1), m.group(2)
            if os.path.exists(path):
                pic = doc.add_paragraph(); pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pic.add_run().add_picture(path, width=Inches(6.6))
                FIG["n"] += 1
                cp = doc.add_paragraph(style="Cap"); runs(cp, f"Figure {FIG['n']}. {alt}")
            i += 1; continue

        if re.match(r"^(---|\*\*\*|___)\s*$", s): i += 1; continue

        m = re.match(r"^(\s*)([-*+]|\d+\.)\s+(.*)$", line)
        if m:
            ordered = bool(re.match(r"\d+\.", m.group(2)))
            p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
            runs(p, m.group(3))
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            i += 1; continue

        if not s: i += 1; continue

        buf = [s]; i += 1
        while i < len(lines):
            nx = lines[i].strip()
            if (not nx or nx.startswith(("#", ">", "|", "```", "---"))
                    or re.match(r"^(\s*)([-*+]|\d+\.)\s+", lines[i])): break
            buf.append(nx); i += 1
        p = doc.add_paragraph(); runs(p, " ".join(buf))

# ------------------------------------------------------------- images

def pick(folder):
    d = os.path.join("08_Screenshots", folder)
    if not os.path.isdir(d): return []
    allf = sorted(os.listdir(d))
    out = []
    for pat in SHOTS.get(folder, []):
        for f in allf:
            if f.lower().endswith(".png") and f.startswith(pat) or (pat in f and f.lower().endswith(".png")):
                if f not in out: out.append(f); break
    return [os.path.join(d, f) for f in out]

def cap_for(fname):
    base = os.path.basename(fname)
    m = re.match(r"^(\d+\.\d+[a-z]?)", base)
    if m:
        key = m.group(1)
        txt = CAPTIONS.get(key) or CAPTIONS.get(key.rstrip("b"))
        return key, txt or re.sub(r"^[\d.]+[a-z]?-", "", os.path.splitext(base)[0]).replace("-", " ")
    key = os.path.splitext(base)[0]
    return None, CAPTIONS.get(key, key.replace("-", " "))

def image_grid(doc, paths, title):
    if not paths: return 0
    h = doc.add_paragraph(style="Heading 4"); runs(h, title)
    t = doc.add_table(rows=0, cols=2); no_borders(t)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx in range(0, len(paths), 2):
        row = t.add_row().cells
        for j in (0, 1):
            if idx + j >= len(paths):
                continue
            p = paths[idx + j]
            cell = row[j]
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                cell.paragraphs[0].add_run().add_picture(p, width=Inches(IMG_W))
            except Exception as e:
                cell.paragraphs[0].add_run(f"[{os.path.basename(p)}]")
            FIG["n"] += 1
            key, txt = cap_for(p)
            cp = cell.add_paragraph(style="Cap")
            runs(cp, f"Fig {FIG['n']}" + (f", shot {key}" if key else "") + f". {txt}")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return len(paths)

# ------------------------------------------------------------- scripts

def commands_only(path):
    """Strip comment lines and blanks, keep the commands."""
    out, blank = [], False
    for ln in open(path, encoding="utf-8").read().split("\n"):
        t = ln.rstrip()
        if t.strip().startswith("!") or not t.strip():
            blank = True; continue
        if blank and out: out.append("")
        blank = False
        out.append(t)
    return out

# ------------------------------------------------------------- build

def read(p): return open(p, encoding="utf-8").read()

def pbreak(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def main():
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Inches(0.55)
        s.top_margin = s.bottom_margin = Inches(0.5)
    setup(doc)

    render(doc, read(f"{SRC}/c1-front.md"))
    render(doc, read(f"{SRC}/c2-threats.md"))
    render(doc, read(f"{SRC}/c3-plan.md"))
    render(doc, read(f"{SRC}/c4-tables.md"))
    render(doc, read(f"{SRC}/c5-configs.md"))

    # screenshots
    p = doc.add_paragraph(style="Heading 2"); runs(p, "7.5 Screenshot evidence")
    p = doc.add_paragraph(style="Note")
    runs(p, "A representative selection. The full set of 90 screenshots, covering every "
            "verification step, accompanies the submission.")
    total = 0
    for _, cfgs in GROUPS:
        for folder, title in cfgs:
            total += image_grid(doc, pick(folder), title)

    render(doc, read(f"{SRC}/c6-evaluation.md"))
    render(doc, read(f"{SRC}/c7-topology-appendix.md"))

    # appendix B
    pbreak(doc)
    p = doc.add_paragraph(style="Heading 1"); runs(p, "Appendix B. Configuration Commands")
    p = doc.add_paragraph(style="Note")
    runs(p, "Commands exactly as entered, with explanatory comments removed for length. "
            "The fully commented scripts, including every Packet Tracer substitution and its "
            "production equivalent, accompany the submission.")
    for folder, items in [
        ("01_Member_1", [("cfg-1-aaa-tacacs-radius.txt", "B1. Configuration 1, AAA"),
                         ("cfg-2-ssh-mgmt-hardening.txt", "B2. Configuration 2, SSH and management plane")]),
        ("02_Member_2", [("cfg-3-vlan-segmentation-acl.txt", "B3. Configuration 3, VLANs and ACLs"),
                         ("cfg-4-port-security.txt", "B4. Configuration 4, port security")]),
        ("03_Member_3", [("cfg-5-edge-firewall-nat.txt", "B5. Configuration 5, edge firewall"),
                         ("cfg-6-dhcp-snooping-dai.txt", "B6. Configuration 6, DHCP snooping and DAI")]),
        ("04_Member_4", [("cfg-7-netflow-export.txt", "B7. Configuration 7, NetFlow"),
                         ("cfg-8-syslog-ntp.txt", "B8. Configuration 8, syslog and NTP")]),
    ]:
        for fn, title in items:
            path = os.path.join(folder, "configs", fn)
            if not os.path.exists(path): continue
            h = doc.add_paragraph(style="Heading 3"); runs(h, title)
            for ln in commands_only(path):
                par = doc.add_paragraph(style="Code"); par.add_run(ln if ln.strip() else " ")

    doc.save(OUT)
    print(f"  {OUT}")
    print(f"  figures: {FIG['n']}   (of 90 captured)")

if __name__ == "__main__":
    main()
